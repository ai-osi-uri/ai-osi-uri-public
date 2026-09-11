"""
議事録 docx 生成スクリプト — config 駆動

使い方:
    python build_minutes.py <minutes.json> <output.docx>

minutes.json のスキーマは references/data-schema.md を参照。
Claude が文字起こしから構造化した JSON を入力として受け取り、
シンプルで読みやすい docx を生成する。
"""
import json
import sys
from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH


# ===== Style helpers =====
ACCENT_RED = RGBColor(0xB9, 0x1C, 0x1C)
TEXT_DARK = RGBColor(0x11, 0x18, 0x27)
TEXT_MED = RGBColor(0x4B, 0x55, 0x63)
TEXT_LIGHT = RGBColor(0x9C, 0xA3, 0xAF)


def set_run(run, *, size=11, bold=False, color=None, italic=False, font="游ゴシック"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    if color is not None:
        run.font.color.rgb = color
    # Set East Asian font
    rPr = run._element.get_or_add_rPr()
    from docx.oxml.ns import qn
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from lxml import etree
        rFonts = etree.SubElement(rPr, qn("w:rFonts"))
    rFonts.set(qn("w:eastAsia"), font)


def add_heading(doc, text, *, size=18, color=ACCENT_RED, space_before=12, space_after=6):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    run = p.add_run(text)
    set_run(run, size=size, bold=True, color=color)
    return p


def add_subheading(doc, text):
    return add_heading(doc, text, size=13, color=TEXT_DARK,
                       space_before=10, space_after=4)


def add_paragraph(doc, text, *, size=10.5, color=TEXT_DARK, indent=0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Cm(indent)
    p.paragraph_format.space_after = Pt(3)
    run = p.add_run(text)
    set_run(run, size=size, color=color)
    return p


def add_bullet(doc, text, *, level=0, size=10.5, color=TEXT_DARK):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.left_indent = Cm(0.5 + level * 0.5)
    p.paragraph_format.space_after = Pt(2)
    run = p.add_run(text)
    set_run(run, size=size, color=color)
    return p


def add_quote(doc, text, *, source=None):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.0)
    p.paragraph_format.right_indent = Cm(1.0)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(f"「{text}」")
    set_run(run, size=10.5, italic=True, color=TEXT_DARK)
    if source:
        ps = doc.add_paragraph()
        ps.paragraph_format.left_indent = Cm(1.0)
        ps.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        ps.paragraph_format.space_after = Pt(8)
        run = ps.add_run(f"── {source}")
        set_run(run, size=9, color=TEXT_MED, italic=True)


def add_action_table(doc, actions):
    """次アクション表: TODO / 期限 / 担当"""
    if not actions:
        return
    table = doc.add_table(rows=1, cols=3)
    table.autofit = False
    table.columns[0].width = Cm(9)
    table.columns[1].width = Cm(3)
    table.columns[2].width = Cm(3)
    # Header
    hdr = table.rows[0].cells
    for i, label in enumerate(["TODO", "期限", "担当"]):
        run = hdr[i].paragraphs[0].add_run(label)
        set_run(run, size=10, bold=True, color=ACCENT_RED)
    # Rows
    for a in actions:
        row = table.add_row().cells
        for col_idx, key in enumerate(["todo", "due", "owner"]):
            val = a.get(key, "—")
            run = row[col_idx].paragraphs[0].add_run(str(val))
            set_run(run, size=10, color=TEXT_DARK)


# ===== Main builder =====
def build_minutes(cfg, output_path):
    doc = Document()
    # Page margins
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    meta = cfg.get("meta", {})

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_after = Pt(4)
    run = title_p.add_run(f"{meta.get('company_name', '')} 商談議事録")
    set_run(run, size=22, bold=True, color=TEXT_DARK)

    # Sub-meta
    subline = []
    if meta.get("date"):
        subline.append(meta["date"])
    if meta.get("attendees"):
        subline.append("出席：" + " / ".join(meta["attendees"]))
    if meta.get("location"):
        subline.append(f"場所：{meta['location']}")
    if subline:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run("　|　".join(subline))
        set_run(run, size=10, color=TEXT_MED, italic=True)

    # Topic
    if meta.get("topic"):
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(12)
        run = p.add_run(f"議題：{meta['topic']}")
        set_run(run, size=11, bold=True, color=TEXT_DARK)

    # サマリ
    summary = cfg.get("summary", [])
    if summary:
        add_heading(doc, "サマリ")
        for s in summary:
            add_bullet(doc, s)

    # 次アクション
    actions = cfg.get("next_actions", [])
    if actions:
        add_heading(doc, "次アクション")
        add_action_table(doc, actions)

    # 相手の重要発言
    quotes = cfg.get("quotes", [])
    if quotes:
        add_heading(doc, "相手の重要発言")
        for q in quotes:
            text = q.get("text", q) if isinstance(q, dict) else q
            source = q.get("source") if isinstance(q, dict) else None
            add_quote(doc, text, source=source)

    # 案件ステータスへの示唆
    status_hint = cfg.get("status_hint")
    if status_hint:
        add_heading(doc, "案件ステータスへの示唆")
        for sh in (status_hint if isinstance(status_hint, list) else [status_hint]):
            add_bullet(doc, sh)

    # 補足メモ
    notes = cfg.get("notes")
    if notes:
        add_heading(doc, "補足メモ")
        for n in (notes if isinstance(notes, list) else [notes]):
            add_paragraph(doc, n)

    # Footer
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    run = p.add_run("──")
    set_run(run, size=8, color=TEXT_LIGHT)
    p2 = doc.add_paragraph()
    run = p2.add_run(
        f"{meta.get('company_name', '')} × AI OSI URI　"
        f"|　生成: {datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    set_run(run, size=8, color=TEXT_LIGHT, italic=True)

    doc.save(output_path)
    return output_path


def main():
    if len(sys.argv) != 3:
        print("Usage: python build_minutes.py <minutes.json> <output.docx>")
        sys.exit(2)
    config_path, output_path = sys.argv[1], sys.argv[2]
    with open(config_path) as f:
        cfg = json.load(f)
    build_minutes(cfg, output_path)
    print(f"WROTE {output_path}")


if __name__ == "__main__":
    main()
