#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""営業管理表（xlsx）を読み書きする唯一の入口。

osi-sales の各スキルは、営業管理表を openpyxl で直接触らず、必ずこのスクリプトを経由する。
理由は「ヘッダが5行目」「行を動かさない」「凡例ブロックの上に挿入する」といった
台帳固有の作法を1箇所に閉じ込めるため。各スキルが個別に openpyxl を書くと必ずズレる。

不変条件（破ったら台帳が壊れる）:
  1. 行は消さない・動かさない。1案件1行、リードから失注まで同じ行に残す。
  2. ヘッダは5行目、データは6行目以降。凡例ブロックより下には案件行を置かない。
  3. ステータスは STATUSES の10値のみ。KPI集計が壊れるので値を増やさない。
  4. 顧客課題メモは追記であって上書きではない（update は append_memo を使う）。
  5. 書き込み前に必ずバックアップを取る（--no-backup を明示しない限り自動で取る）。

使い方:
  ledger.py read [--status S] [--owner O] [--case-id ID]
  ledger.py find --company NAME | --case-id ID
  ledger.py next-id
  ledger.py append --data '{"取引先名":"...", "ステータス":"リード", ...}'
  ledger.py update --row N   --data '{"ステータス":"商談中"}' [--append-memo "9/15 初回商談"]
  ledger.py update --case-id 139 --data '{...}'
  ledger.py columns
"""
import argparse, json, os, re, shutil, sys, unicodedata
from datetime import datetime

try:
    import openpyxl
    from copy import copy
except ImportError:
    sys.exit("openpyxl が要る: pip install openpyxl --break-system-packages")

# --- 組織固有値は osi-profile.md（会社プロファイル）から読む。環境変数で上書きできる ---
# 探索順: $OSI_PROFILE → cwd の祖先の osi-profile.md / _shared/osi-profile.md
#         → Cowork マウント ($HOME/mnt/*/osi-profile.md) → Google Drive 同期フォルダ
# 雛形: ai-osi-uri-plugins/config/osi-profile.example.md（無ければ質問して作る）

def _parse_frontmatter(text):
    m = re.match(r"^---\s*\n(.*?)\n---", text, re.S)
    if not m:
        return {}
    try:
        import yaml  # type: ignore
        return yaml.safe_load(m.group(1)) or {}
    except ImportError:
        # 最小パーサ: 2 階層の "key: value" だけ読む（PyYAML が無い環境向け）
        d, cur = {}, None
        for line in m.group(1).splitlines():
            if not line.strip() or line.lstrip().startswith("#"):
                continue
            k, _, v = line.partition(":")
            v = v.split("#")[0].strip().strip("\"'")
            if not line.startswith(" "):
                cur = k.strip(); d[cur] = {} if v == "" else v
            elif isinstance(d.get(cur), dict):
                d[cur][k.strip()] = v
        return d


def _find_profile():
    import glob
    env = os.environ.get("OSI_PROFILE")
    if env and os.path.exists(env):
        return env
    d = os.getcwd()
    while True:
        for c in (os.path.join(d, "osi-profile.md"), os.path.join(d, "_shared", "osi-profile.md")):
            if os.path.exists(c):
                return c
        nd = os.path.dirname(d)
        if nd == d:
            break
        d = nd
    pats = [os.path.expanduser("~/mnt/*/osi-profile.md"),
            os.path.expanduser("~/Library/CloudStorage/GoogleDrive-*/共有ドライブ/*/osi-profile.md"),
            os.path.expanduser("~/Library/CloudStorage/GoogleDrive-*/マイドライブ/*/osi-profile.md")]
    for pat in pats:
        hits = sorted(glob.glob(pat))
        if hits:
            return hits[0]
    return None


PROFILE_PATH = _find_profile()
PROFILE = _parse_frontmatter(open(PROFILE_PATH, encoding="utf-8").read()) if PROFILE_PATH else {}
_paths = PROFILE.get("paths", {}) or {}
_ledgers = PROFILE.get("ledgers", {}) or {}

if PROFILE_PATH:
    _root_dir = os.path.dirname(PROFILE_PATH)
    if os.path.basename(_root_dir) == "_shared":
        _root_dir = os.path.dirname(_root_dir)
    _root_dir = os.path.normpath(os.path.join(_root_dir, str(_paths.get("root", "."))))
else:
    _root_dir = None
DRIVE_ROOT = os.environ.get("OSI_DRIVE_ROOT", _root_dir or "")
LEDGER = os.environ.get("OSI_SALES_LEDGER",
                        os.path.join(DRIVE_ROOT, _ledgers.get("sales", "営業管理/営業管理表.xlsx")))
CASE_ROOT = os.environ.get("OSI_CASE_ROOT", os.path.join(DRIVE_ROOT, _paths.get("projects", "案件")))
TAB = os.environ.get("OSI_SALES_TAB", _ledgers.get("sales_tab", "取引先管理"))
HEADER_ROW = int(os.environ.get("OSI_SALES_HEADER_ROW", _ledgers.get("sales_header_row", 5)))
FIRST_DATA_ROW = HEADER_ROW + 1

STATUSES = ["リード", "商談中", "提案中", "金額提示済み", "受注",
            "契約済み", "デリバリー中", "保留", "失注", "終了"]
SUBFOLDERS = ["01_提案・見積", "02_契約", "03_制作・成果物", "04_請求", "05_受領資料"]


def _norm(s):
    """表記ゆれ吸収: 全半角・大文字小文字・法人格・記号・空白を落とす。"""
    if s is None:
        return ""
    s = unicodedata.normalize("NFKC", str(s)).lower()
    for w in ("株式会社", "有限会社", "合同会社", "一般社団法人", "医療法人", "(株)", "(有)", "様"):
        s = s.replace(w, "")
    return re.sub(r"[\s\-‐―ー・/／,，.．()（）&＆'\"]", "", s)


def _open():
    if not PROFILE_PATH and not os.environ.get("OSI_SALES_LEDGER"):
        sys.exit("会社プロファイル osi-profile.md が見つからない。連結フォルダ直下に置く"
                 "（雛形: ai-osi-uri-plugins/config/osi-profile.example.md）か、OSI_PROFILE で場所を指定する")
    if not os.path.exists(LEDGER):
        sys.exit(f"台帳が無い: {LEDGER}（osi-profile.md の ledgers.sales を確認）")
    lock = os.path.join(os.path.dirname(LEDGER), "~$" + os.path.basename(LEDGER))
    if os.path.exists(lock):
        sys.exit(f"Excelで開かれたままのロック残骸がある: {lock}\n"
                 "Finder で削除してもらってから再実行する（bash からは消せないことがある）。")
    wb = openpyxl.load_workbook(LEDGER)
    return wb, wb[TAB]


def _cols(ws):
    return {ws.cell(HEADER_ROW, c).value: c
            for c in range(1, ws.max_column + 1) if ws.cell(HEADER_ROW, c).value}


def _last_case_row(ws, H):
    """案件行の最終行。凡例など注記ブロックは A列にしか値が無いので除外する。"""
    last = FIRST_DATA_ROW - 1
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        if ws.cell(r, H["取引先名"]).value:
            last = r
    return last


def _rows(ws, H):
    out = []
    for r in range(FIRST_DATA_ROW, _last_case_row(ws, H) + 1):
        if not ws.cell(r, H["取引先名"]).value:
            continue
        d = {"__row": r}
        for name, c in H.items():
            v = ws.cell(r, c).value
            if v not in (None, ""):
                d[name] = v if not isinstance(v, datetime) else v.strftime("%Y-%m-%d")
        out.append(d)
    return out


def _backup():
    bdir = os.path.join(os.path.dirname(LEDGER), "_backup")
    os.makedirs(bdir, exist_ok=True)
    dst = os.path.join(bdir, f"{datetime.now():%Y%m%d_%H%M%S}_" + os.path.basename(LEDGER))
    shutil.copy2(LEDGER, dst)
    return dst


def _style(ws, src_row, dst_row, col):
    s, t = ws.cell(src_row, col), ws.cell(dst_row, col)
    t.font, t.border = copy(s.font), copy(s.border)
    t.fill, t.alignment = copy(s.fill), copy(s.alignment)
    t.number_format = s.number_format


def cmd_read(a):
    wb, ws = _open(); H = _cols(ws)
    rows = _rows(ws, H)
    if a.status:
        rows = [r for r in rows if r.get("ステータス") == a.status]
    if a.owner:
        rows = [r for r in rows if a.owner in str(r.get("営業担当", ""))]
    if a.case_id:
        rows = [r for r in rows if str(r.get("案件ID", "")).zfill(3) == a.case_id.zfill(3)]
    print(json.dumps(rows, ensure_ascii=False, indent=2))


def cmd_find(a):
    wb, ws = _open(); H = _cols(ws)
    rows = _rows(ws, H)
    if a.case_id:
        hit = [r for r in rows if str(r.get("案件ID", "")).zfill(3) == a.case_id.zfill(3)]
    else:
        q = _norm(a.company)
        hit = [r for r in rows if q and (q in _norm(r.get("取引先名")) or _norm(r.get("取引先名")) in q)]
    print(json.dumps(hit, ensure_ascii=False, indent=2))


def cmd_next_id(a):
    """採番の正本は Drive のフォルダ名。台帳のZ列ではない。"""
    ids = []
    if os.path.isdir(CASE_ROOT):
        for n in os.listdir(CASE_ROOT):
            m = re.match(r"^(\d{3})\.", n)
            if m:
                ids.append(int(m.group(1)))
    if not ids:
        sys.exit(f"案件フォルダが1つも見えない: {CASE_ROOT}\n"
                 "Drive のキャッシュが外れている可能性がある。「無い」と報告せず、"
                 "Finder で開いてもらうか Drive コネクタで一覧を取ること。")
    print(json.dumps({"case_root": CASE_ROOT, "max": max(ids),
                      "next": f"{max(ids) + 1:03d}", "count": len(ids)}, ensure_ascii=False))


def cmd_append(a):
    data = json.loads(a.data)
    if not data.get("取引先名"):
        sys.exit("取引先名は必須")
    st = data.get("ステータス")
    if st and st not in STATUSES:
        sys.exit(f"ステータスは {STATUSES} のいずれか（KPI集計が壊れる）: {st}")
    wb, ws = _open(); H = _cols(ws)
    unknown = [k for k in data if k not in H]
    if unknown:
        sys.exit(f"存在しない列: {unknown}\n使える列: {list(H)}")
    dup = [r for r in _rows(ws, H) if _norm(r.get("取引先名")) == _norm(data["取引先名"])]
    if dup and not a.allow_duplicate:
        sys.exit("同名の行が既にある。既存行を update するか、別案件なら --allow-duplicate:\n"
                 + json.dumps(dup, ensure_ascii=False, indent=2))
    if not a.no_backup:
        print("backup:", _backup(), file=sys.stderr)
    last = _last_case_row(ws, H)
    ws.insert_rows(last + 1, 1)
    for k, v in data.items():
        ws.cell(last + 1, H[k]).value = v
        _style(ws, last, last + 1, H[k])
    ws.row_dimensions[last + 1].height = ws.row_dimensions[last].height
    wb.save(LEDGER)
    print(json.dumps({"inserted_row": last + 1, "案件ID": data.get("案件ID"),
                      "取引先名": data["取引先名"]}, ensure_ascii=False))


def cmd_update(a):
    data = json.loads(a.data) if a.data else {}
    st = data.get("ステータス")
    if st and st not in STATUSES:
        sys.exit(f"ステータスは {STATUSES} のいずれか: {st}")
    wb, ws = _open(); H = _cols(ws)
    if a.row:
        row = a.row
    else:
        hit = [r for r in _rows(ws, H) if str(r.get("案件ID", "")).zfill(3) == a.case_id.zfill(3)]
        if len(hit) != 1:
            sys.exit("案件IDが一意に引けない。--row で行番号を指定する:\n"
                     + json.dumps(hit, ensure_ascii=False, indent=2))
        row = hit[0]["__row"]
    unknown = [k for k in data if k not in H]
    if unknown:
        sys.exit(f"存在しない列: {unknown}")
    if not a.no_backup:
        print("backup:", _backup(), file=sys.stderr)
    for k, v in data.items():
        ws.cell(row, H[k]).value = v
    if a.append_memo:
        c = H["顧客課題メモ"]
        cur = ws.cell(row, c).value
        ws.cell(row, c).value = (str(cur).rstrip() + " / " + a.append_memo) if cur else a.append_memo
    wb.save(LEDGER)
    print(json.dumps({"updated_row": row, "取引先名": ws.cell(row, H["取引先名"]).value},
                     ensure_ascii=False))


def cmd_columns(a):
    wb, ws = _open()
    print(json.dumps({"profile": PROFILE_PATH, "ledger": LEDGER, "tab": TAB, "header_row": HEADER_ROW,
                      "columns": list(_cols(ws)), "statuses": STATUSES,
                      "case_root": CASE_ROOT, "subfolders": SUBFOLDERS},
                     ensure_ascii=False, indent=2))


def main():
    p = argparse.ArgumentParser(description="営業管理表(xlsx)の読み書き")
    sub = p.add_subparsers(dest="cmd", required=True)
    for name, fn in (("read", cmd_read), ("find", cmd_find), ("next-id", cmd_next_id),
                     ("append", cmd_append), ("update", cmd_update), ("columns", cmd_columns)):
        s = sub.add_parser(name)
        s.set_defaults(func=fn)
        if name in ("read", "find", "update"):
            s.add_argument("--case-id")
        if name in ("read",):
            s.add_argument("--status"); s.add_argument("--owner")
        if name in ("find",):
            s.add_argument("--company")
        if name in ("append", "update"):
            s.add_argument("--data"); s.add_argument("--no-backup", action="store_true")
        if name == "append":
            s.add_argument("--allow-duplicate", action="store_true")
        if name == "update":
            s.add_argument("--row", type=int); s.add_argument("--append-memo")
    a = p.parse_args()
    a.func(a)


if __name__ == "__main__":
    main()
