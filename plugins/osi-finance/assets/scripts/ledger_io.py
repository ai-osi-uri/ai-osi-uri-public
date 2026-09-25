#!/usr/bin/env python3
"""台帳（ローカル xlsx）の読み書き — AI OSI URI Finance 拡張（.mcpb）が無い環境で sheets_* の代わりに使う

拡張の道具があれば拡張を使う（コンソールと同じ経路で書ける）。**拡張が無いときはこれを使う。**
引数と戻りは拡張の道具にそろえてある（タブ名だけで引く・列名→値で書く・dedupe_by・dry_run・path_fixes）。

使い方（<DIR> は経理フォルダ＝台帳 xlsx を置いたフォルダ。出力はすべて JSON）:
    python3 ledger_io.py <DIR> tabs                                            # ≒ sheets_list_tabs
    python3 ledger_io.py <DIR> read --tab 月次請求スケジュール [--where 請求ステータス=未請求 ...]
                                                                               # ≒ sheets_read_schedule / sheets_get_values
    python3 ledger_io.py <DIR> append --tab 月次支払管理 --row '{"支払先ID":"V-001", ...}' \\
            [--dedupe-by 請求書番号 支払先ID] [--dry-run] [--allow-missing-file]  # ≒ sheets_append_row（--row は配列も可）
    python3 ledger_io.py <DIR> update --tab 月次請求スケジュール --where 契約ID=AR-C-001 --where 対象月=2026-08 \\
            --set '{"請求書番号":"INV-2026-09-001"}' [--dry-run]              # 列名で指定して値を更新（≒ sheets_update_values）
    python3 ledger_io.py <DIR> update-status --updates '[{"inv":"INV-2026-09-001","status":"請求済"}]'
                                                                               # ≒ sheets_update_status（合算請求は全行）
    python3 ledger_io.py <DIR> next-id --tab 取引先マスタ --col 取引先ID --prefix P- [--width 3]
    python3 ledger_io.py <DIR> expand-schedule --contract AR-C-001 [--dry-run] [--bill-same-month] [--month YYYY-MM]
                                                                               # ≒ ledger_maintain expand_schedule

--row / --set / --updates は JSON 文字列。先頭を @ にするとファイルから読む（例 --row @row.json）。

守ること（data-layout.yaml の「読み書きの原則」と同じ）
------------------------------------------------------
* 列はヘッダー名で解決する。ヘッダー行の位置は data-layout.yaml の header_row（無ければ先頭10行で推定）。
* 追記位置は「末尾から遡って最初の非空行の次」。件数から計算しない（途中の空行で既存行を潰した事故がある）。
* 台帳に無い列名はエラー（黙って捨てない）。
* 書く前に <DIR>/_backup/ へ控えを取る。Excel で開かれている（~$ ロックファイルがある）ときは書かない。
* 作成者/作成日時（新規行）・更新者/更新日時（更新）は自動で押す。操作者は --actor → 環境変数
  OSI_FINANCE_OPERATOR → settings の OPERATOR_EMAIL の順。どれも無ければ空で書く（推測で埋めない）。
* 「〜ファイル」列は台帳フォルダからのフォルダ込み相対パス。実ファイルが無ければエラー（同名が1つだけ見つかれば
  そこへ直して path_fixes で報告）。--allow-missing-file で先に起票できる。
* 請求ステータスは下げない（差し戻しは人の操作のみ）。update-status は下げる更新を拒否する。

終了コード: 0=成功 / 1=エラー（書いていない） / 2=対象行が無い
"""
import argparse
import calendar
import datetime as dt
import json
import os
import re
import shutil
import sys
from pathlib import Path

try:
    import openpyxl
    import yaml
except ImportError as e:
    sys.exit(f"openpyxl と PyYAML が必要です: pip3 install openpyxl pyyaml ({e})")

HERE = Path(__file__).resolve().parent
SCHEMA = HERE.parent / "schema" / "data-layout.yaml"
SETTINGS_NAME = "osi-finance-settings.md"
TEXT_COL_HINTS = ("ID", "番号", "対象月", "発生月", "発生日", "郵便番号", "コード", "旧ID", "envelope_id")
STATUS_ORDER = ["未請求", "下書き済", "請求済", "入金済"]
DEFAULT_PATTERNS = {"月額固定": "monthly", "一括": "lump", "都度": "variable"}


class LedgerError(Exception):
    pass


# ---------------------------------------------------------------- 仕様・設定

def load_schema():
    try:
        return yaml.safe_load(SCHEMA.read_text(encoding="utf-8"))
    except OSError:
        return {"ledgers": {}}


def tab_spec(schema, tab):
    for lg in (schema.get("ledgers") or {}).values():
        spec = (lg.get("tabs") or {}).get(tab)
        if spec is not None:
            return spec
    return {}


def col_spec(spec, name):
    return next((c for c in spec.get("columns", []) if c.get("name") == name), {})


def operator(args_actor, d):
    if args_actor:
        return args_actor
    if os.environ.get("OSI_FINANCE_OPERATOR"):
        return os.environ["OSI_FINANCE_OPERATOR"]
    p = d / SETTINGS_NAME
    if p.exists():
        for line in p.read_text(encoding="utf-8").splitlines():
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if len(cells) >= 2 and cells[0] == "OPERATOR_EMAIL" and cells[1] and "{{" not in cells[1]:
                return cells[1]
    return ""


def now():
    return dt.datetime.now().strftime("%Y-%m-%d %H:%M")


def s(v):
    if v is None:
        return ""
    if isinstance(v, dt.datetime):
        return v.strftime("%Y-%m-%d") if (v.hour, v.minute, v.second) == (0, 0, 0) else v.strftime("%Y-%m-%d %H:%M")
    if isinstance(v, dt.date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


# ---------------------------------------------------------------- 台帳の場所

def ledger_files(d):
    return sorted(p for p in d.glob("*.xlsx") if not p.name.startswith(("~$", ".")))


def find_tab(d, tab):
    hits = []
    for p in ledger_files(d):
        try:
            wb = openpyxl.load_workbook(p, read_only=True)
        except Exception:  # noqa: BLE001
            continue
        if tab in wb.sheetnames:
            hits.append(p)
        wb.close()
    if not hits:
        known = sorted({t for p in ledger_files(d) for t in openpyxl.load_workbook(p, read_only=True).sheetnames})
        raise LedgerError(f"タブが見つかりません: {tab}（{d} にあるタブ: {', '.join(known)}）")
    if len(hits) > 1:
        raise LedgerError(f"タブ「{tab}」が複数のファイルにあります: {[p.name for p in hits]}（タブ名は台帳をまたいで一意にする）")
    return hits[0]


def header_row_of(ws, spec):
    hr = spec.get("header_row")
    want = [c["name"] for c in spec.get("columns", [])]
    if hr:
        vals = [s(c.value) for c in ws[int(hr)]] if int(hr) <= ws.max_row else []
        if want and any(v in want for v in vals):
            return int(hr)
    best, n_best = 1, -1
    for r in range(1, min(10, ws.max_row) + 1):
        n = sum(1 for c in ws[r] if s(c.value))
        if n > n_best:
            best, n_best = r, n
    return best


def headers(ws, hr):
    return {s(c.value): c.column for c in ws[hr] if s(c.value)}


def last_data_row(ws, hr, cols):
    for r in range(ws.max_row, hr, -1):
        if any(s(ws.cell(row=r, column=c).value) for c in cols):
            return r
    return hr


class Book:
    def __init__(self, d, tab, write=False):
        self.d, self.tab = d, tab
        self.path = find_tab(d, tab)
        if write and (self.path.parent / f"~${self.path.name}").exists():
            raise LedgerError(f"台帳が Excel で開かれています: {self.path.name}（~$ ロックファイルあり）。閉じてから再実行してください")
        self.wb = openpyxl.load_workbook(self.path)
        self.ws = self.wb[tab]
        self.spec = tab_spec(load_schema(), tab)
        self.hr = header_row_of(self.ws, self.spec)
        self.cols = headers(self.ws, self.hr)
        if not self.cols:
            raise LedgerError(f"「{tab}」の見出し行（{self.hr}行目）が空です")

    def rows(self):
        out = []
        last = last_data_row(self.ws, self.hr, self.cols.values())
        for r in range(self.hr + 1, last + 1):
            d = {name: s(self.ws.cell(row=r, column=c).value) for name, c in self.cols.items()}
            if not any(d.values()):
                continue
            d["__row"] = r
            out.append(d)
        return out

    def save(self):
        bdir = self.d / "_backup"
        bdir.mkdir(parents=True, exist_ok=True)
        bak = bdir / f"{self.path.stem}_{dt.datetime.now():%Y%m%d-%H%M%S-%f}.xlsx"
        shutil.copy2(self.path, bak)
        self.wb.save(self.path)
        return str(bak)

    def put(self, r, name, value):
        c = self.ws.cell(row=r, column=self.cols[name])
        if any(h in name for h in TEXT_COL_HINTS) and value is not None and not isinstance(value, (int, float)):
            c.number_format = "@"
            value = str(value)
        c.value = value


# ---------------------------------------------------------------- 値の検査

def normalize_enum(spec, name, value, warnings):
    cs = col_spec(spec, name)
    if not isinstance(value, str) or not value:
        return value
    acc = cs.get("accept") or {}
    if value in acc:
        value = acc[value]
    vals = cs.get("values")
    if vals and value not in vals:
        warnings.append(f"{name}={value!r} は仕様の値（{' / '.join(vals)}）にありません")
    return value


def fix_file_path(d, name, value, allow_missing, path_fixes):
    if not name.endswith("ファイル") or value in (None, ""):
        return value
    v = str(value).replace("\\", "/").lstrip("./")
    if (d / v).exists() and "/" in v:
        return v
    cands = [p for p in d.rglob(Path(v).name) if "_backup" not in p.parts]
    if len(cands) == 1:
        fixed = str(cands[0].relative_to(d)).replace("\\", "/")
        path_fixes.append({"column": name, "from": value, "to": fixed})
        return fixed
    if allow_missing:
        return v
    why = "ファイル名だけです。台帳フォルダからのフォルダ込み相対パスで渡してください" if "/" not in v else "実ファイルがありません"
    raise LedgerError(f"{name}={value!r}: {why}（先に起票するなら --allow-missing-file）")


def parse_json_arg(v):
    if v is None:
        return None
    if v.startswith("@"):
        v = Path(v[1:]).expanduser().read_text(encoding="utf-8")
    return json.loads(v)


def parse_where(items):
    out = {}
    for it in items or []:
        if "=" not in it:
            raise LedgerError(f"--where は 列名=値 で渡す: {it}")
        k, v = it.split("=", 1)
        out[k.strip()] = v.strip()
    return out


def match(row, where):
    return all(row.get(k, "") == v for k, v in where.items())


# ---------------------------------------------------------------- コマンド

def cmd_tabs(d, a):
    return {"ledger_dir": str(d), "files": {p.name: openpyxl.load_workbook(p, read_only=True).sheetnames for p in ledger_files(d)}}


def cmd_read(d, a):
    b = Book(d, a.tab)
    where = parse_where(a.where)
    for k in where:
        if k not in b.cols:
            raise LedgerError(f"列「{k}」は「{a.tab}」にありません（あるのは: {', '.join(b.cols)}）")
    rows = [r for r in b.rows() if match(r, where)]
    return {"tab": a.tab, "file": b.path.name, "header_row": b.hr, "headers": list(b.cols), "count": len(rows), "rows": rows}


def append_rows(d, tab, rows, dedupe_by=None, dry_run=False, allow_missing=False, actor=None):
    b = Book(d, tab, write=not dry_run)
    op = operator(actor, d)
    existing = b.rows()
    last = last_data_row(b.ws, b.hr, b.cols.values())
    results, path_fixes, warnings = [], [], []
    for row in rows:
        unknown = [k for k in row if k not in b.cols]
        if unknown:
            raise LedgerError(f"「{tab}」に無い列: {unknown}（あるのは: {', '.join(b.cols)}）")
        row = {k: normalize_enum(b.spec, k, v, warnings) for k, v in row.items()}
        row = {k: fix_file_path(d, k, v, allow_missing, path_fixes) for k, v in row.items()}
        if dedupe_by:
            miss = [k for k in dedupe_by if k not in b.cols]
            if miss:
                raise LedgerError(f"dedupe_by の列が「{tab}」にありません: {miss}")
            key = [s(row.get(k)) for k in dedupe_by]
            dup = next((e for e in existing if [e.get(k, "") for k in dedupe_by] == key), None)
            if dup:
                results.append({"appended": False, "duplicate_of": dup["__row"], "key": dict(zip(dedupe_by, key))})
                continue
        last += 1
        full = dict(row)
        for k, v in (("作成者", op), ("作成日時", now())):
            if k in b.cols and not s(full.get(k)):
                full[k] = v
        if not dry_run:
            for k, v in full.items():
                b.put(last, k, v)
        existing.append({**{k: s(v) for k, v in full.items()}, "__row": last})
        results.append({"appended": True, "row": last, "values": {k: s(v) for k, v in full.items()}})
    out = {"tab": tab, "file": b.path.name, "header_row": b.hr, "dry_run": dry_run, "results": results,
           "path_fixes": path_fixes, "warnings": warnings}
    if not dry_run and any(r["appended"] for r in results):
        out["backup"] = b.save()
        # 読み戻して検算する（拡張と同じ）
        chk = {r["__row"]: r for r in Book(d, tab).rows()}
        bad = [r["row"] for r in results if r.get("appended") and r["row"] not in chk]
        if bad:
            raise LedgerError(f"書き込み後の読み戻しで行が見つかりません: {bad}")
    return out


def cmd_append(d, a):
    rows = parse_json_arg(a.row)
    rows = rows if isinstance(rows, list) else [rows]
    return append_rows(d, a.tab, rows, a.dedupe_by, a.dry_run, a.allow_missing_file, a.actor)


def update_rows(d, tab, where, values, dry_run=False, actor=None, row_numbers=None, allow_missing=False):
    b = Book(d, tab, write=not dry_run)
    op = operator(actor, d)
    unknown = [k for k in list(where) + list(values) if k not in b.cols]
    if unknown:
        raise LedgerError(f"「{tab}」に無い列: {unknown}（あるのは: {', '.join(b.cols)}）")
    targets = [r for r in b.rows() if match(r, where) and (not row_numbers or r["__row"] in row_numbers)]
    if not targets:
        return None
    warnings, path_fixes, changes = [], [], []
    values = {k: normalize_enum(b.spec, k, v, warnings) for k, v in values.items()}
    values = {k: fix_file_path(d, k, v, allow_missing, path_fixes) for k, v in values.items()}
    for r in targets:
        diff = {k: {"from": r.get(k, ""), "to": s(v)} for k, v in values.items() if r.get(k, "") != s(v)}
        if not diff:
            continue
        if not dry_run:
            for k, v in values.items():
                b.put(r["__row"], k, v)
            for k, v in (("更新者", op), ("更新日時", now())):
                if k in b.cols:
                    b.put(r["__row"], k, v)
        changes.append({"row": r["__row"], "changes": diff})
    out = {"tab": tab, "file": b.path.name, "dry_run": dry_run, "matched": len(targets), "updated": changes,
           "path_fixes": path_fixes, "warnings": warnings}
    if changes and not dry_run:
        out["backup"] = b.save()
    return out


def cmd_update(d, a):
    rn = set(a.row_number or [])
    out = update_rows(d, a.tab, parse_where(a.where), parse_json_arg(a.set), a.dry_run, a.actor, rn, a.allow_missing_file)
    if out is None:
        return {"tab": a.tab, "matched": 0, "error": "条件に合う行がありません"}, 2
    return out


def cmd_update_status(d, a):
    tab = a.tab
    b = Book(d, tab)
    rows = b.rows()
    results = []
    for u in parse_json_arg(a.updates):
        inv = u["inv"]
        hit = [r for r in rows if r.get("請求書番号") == inv]  # 合算請求は1対多（全行を更新する）
        if not hit:
            results.append({"inv": inv, "ok": False, "problem": "請求書番号が台帳にありません"})
            continue
        vals = {}
        if u.get("status"):
            vals["請求ステータス"] = u["status"]
        if u.get("paid_date"):
            vals["入金日"] = u["paid_date"]
        if u.get("paid_amount") is not None:
            vals["入金額"] = u["paid_amount"]
        if u.get("sent_date") and "送付日" in b.cols:
            vals["送付日"] = u["sent_date"]
        new = vals.get("請求ステータス")
        if new in STATUS_ORDER:
            low = [r["__row"] for r in hit if r.get("請求ステータス") in STATUS_ORDER
                   and STATUS_ORDER.index(r["請求ステータス"]) > STATUS_ORDER.index(new)]
            if low:
                results.append({"inv": inv, "ok": False, "rows": low,
                                "problem": "ステータスを下げる更新はしない（差し戻しは人の操作のみ）"})
                continue
        out = update_rows(d, tab, {"請求書番号": inv}, vals, a.dry_run, a.actor)
        results.append({"inv": inv, "ok": True, "rows": [r["__row"] for r in hit], "updated": out["updated"] if out else []})
    return {"tab": tab, "dry_run": a.dry_run, "results": results}


def next_id(d, tab, col, prefix, width=3):
    b = Book(d, tab)
    if col not in b.cols:
        raise LedgerError(f"列「{col}」は「{tab}」にありません")
    rx = re.compile(r"^" + re.escape(prefix) + r"(\d+)$")
    nums = [int(m.group(1)) for r in b.rows() for m in [rx.match(r.get(col, ""))] if m]
    return f"{prefix}{(max(nums) + 1 if nums else 1):0{width}d}"


def cmd_next_id(d, a):
    return {"tab": a.tab, "col": a.col, "next": next_id(d, a.tab, a.col, a.prefix, a.width)}


# ---------------------------------------------------------------- スケジュール展開

def to_date(v):
    if isinstance(v, dt.datetime):
        return v.date()
    if isinstance(v, dt.date):
        return v
    t = s(v).replace("/", "-")
    m = re.match(r"^(\d{4})-(\d{1,2})-(\d{1,2})", t)
    if not m:
        raise LedgerError(f"日付として読めません: {v!r}")
    return dt.date(int(m.group(1)), int(m.group(2)), int(m.group(3)))


def num(v):
    try:
        return float(s(v).replace(",", "").replace("¥", "").replace("￥", "") or 0)
    except ValueError:
        return 0.0


def add_month(y, m, k=1):
    m2 = m - 1 + k
    return y + m2 // 12, m2 % 12 + 1


def month_end(y, m):
    return dt.date(y, m, calendar.monthrange(y, m)[1])


def billing_patterns(d):
    try:
        rows = Book(d, "発行者設定").rows()
    except LedgerError:
        return dict(DEFAULT_PATTERNS)
    v = next((r.get("値", "") for r in rows if r.get("項目") == "金額区分"), "")
    pats = {}
    for part in re.split(r"[;；]", v):
        if "=" in part:
            k, p = part.split("=", 1)
            pats[k.strip()] = p.strip()
    return pats or dict(DEFAULT_PATTERNS)


def cmd_expand(d, a):
    contracts = [r for r in Book(d, "契約マスタ").rows() if r.get("契約ID") == a.contract]
    if not contracts:
        return {"error": f"契約マスタに {a.contract} がありません"}, 2
    c = contracts[0]
    if c.get("状態") != "締結済":
        raise LedgerError(f"{a.contract} は 状態={c.get('状態') or '空'}。展開は 締結済 の契約だけ")
    label = c.get("金額区分", "")
    pat = billing_patterns(d).get(label)
    if not pat:
        raise LedgerError(f"金額区分「{label}」が発行者設定の金額区分に無い（settings §4-4 に足すか既定ラベルに寄せる）")
    if pat == "variable":
        return {"contract": a.contract, "pattern": pat, "rows": [], "note": "都度（variable）は展開しない。実績が出た月に起票する"}
    direction = c.get("方向") or "受注"
    tab = "月次請求スケジュール" if direction == "受注" else "月次支払スケジュール"
    subject_base = c.get("契約内容") or c.get("契約種別") or "業務委託"
    rows = []

    def ar_row(y, m, amount, note=""):
        target = f"{y:04d}-{m:02d}"
        if a.bill_same_month:
            bill = month_end(y, m)
            due = month_end(*add_month(y, m))
        else:
            by, bm = add_month(y, m)
            bill = dt.date(by, bm, 1)
            due = month_end(by, bm)
        row = {"契約ID": a.contract, "取引先": c.get("取引先", ""), "対象月": target,
               "件名": f"{subject_base}（{y}年{m}月分）", "請求額(税込)": int(round(amount)),
               "請求日": bill.isoformat(), "支払期限": due.isoformat(), "請求ステータス": "未請求"}
        if note:
            row["備考"] = note
        return row

    if pat == "monthly":
        start, end = to_date(c.get("契約開始")), to_date(c.get("契約終了"))
        monthly = num(c.get("月額(税込)"))
        if monthly <= 0:
            raise LedgerError(f"{a.contract} の 月額(税込) が空か0")
        y, m = start.year, start.month
        while (y, m) <= (end.year, end.month):
            days = calendar.monthrange(y, m)[1]
            d0 = start.day if (y, m) == (start.year, start.month) else 1
            d1 = end.day if (y, m) == (end.year, end.month) else days
            n = d1 - d0 + 1
            amount, note = monthly, ""
            if n < days:
                amount = monthly * n / days
                note = f"日割（{n}/{days}日）。契約書の日割規定を確認"
            rows.append((y, m, amount, note))
            y, m = add_month(y, m)
    elif pat in ("lump", "prepaid"):
        if not a.month:
            raise LedgerError(f"金額区分={label}（{pat}）は該当月を人が決める。--month YYYY-MM で指定する")
        y, m = map(int, a.month.split("-"))
        amount = num(c.get("一括額(税込)")) or num(c.get("数量")) * num(c.get("単価(税込)"))
        rows.append((y, m, amount, ""))
    else:
        raise LedgerError(f"未対応の課金パターン: {pat}")

    if tab == "月次請求スケジュール":
        out_rows = [ar_row(*r) for r in rows]
    else:
        payee = ""
        try:
            payee = next((p.get("支払先ID", "") for p in Book(d, "支払先マスタ").rows()
                          if p.get("取引先ID") == c.get("取引先ID")), "")
        except LedgerError:
            pass
        out_rows = []
        for y, m, amount, note in rows:
            py, pm = add_month(y, m)
            out_rows.append({"予定ID": f"{payee or c.get('取引先ID', '')}-{y:04d}{m:02d}", "契約ID": a.contract,
                             "取引先ID": c.get("取引先ID", ""), "支払先ID": payee, "支払先名": c.get("取引先", ""),
                             "対象月": f"{y:04d}-{m:02d}", "予定額(税込)": int(round(amount)),
                             "支払予定日": month_end(py, pm).isoformat(), "状態": "未払予定"})
    res = append_rows(d, tab, out_rows, ["契約ID", "対象月"], a.dry_run, False, a.actor)
    res.update({"contract": a.contract, "pattern": pat})
    return res


# ---------------------------------------------------------------- main

def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("ledger_dir", help="経理フォルダ（台帳 xlsx を置いたフォルダ）")
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("tabs")
    p = sub.add_parser("read")
    p.add_argument("--tab", required=True)
    p.add_argument("--where", action="append")
    p = sub.add_parser("append")
    p.add_argument("--tab", required=True)
    p.add_argument("--row", required=True)
    p.add_argument("--dedupe-by", nargs="+")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--allow-missing-file", action="store_true")
    p.add_argument("--actor")
    p = sub.add_parser("update")
    p.add_argument("--tab", required=True)
    p.add_argument("--where", action="append")
    p.add_argument("--row-number", type=int, action="append", help="__row で行を限定する（read の戻りの __row）")
    p.add_argument("--set", required=True)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--allow-missing-file", action="store_true")
    p.add_argument("--actor")
    p = sub.add_parser("update-status")
    p.add_argument("--tab", default="月次請求スケジュール")
    p.add_argument("--updates", required=True)
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--actor")
    p = sub.add_parser("next-id")
    p.add_argument("--tab", required=True)
    p.add_argument("--col", required=True)
    p.add_argument("--prefix", required=True)
    p.add_argument("--width", type=int, default=3)
    p = sub.add_parser("expand-schedule")
    p.add_argument("--contract", required=True)
    p.add_argument("--month", help="lump / prepaid の該当月 YYYY-MM")
    p.add_argument("--bill-same-month", action="store_true", help="請求日=対象月末（契約で当月請求と定めたもの）")
    p.add_argument("--dry-run", action="store_true")
    p.add_argument("--actor")
    a = ap.parse_args()
    d = Path(a.ledger_dir).expanduser()
    if not d.is_dir():
        print(json.dumps({"error": f"経理フォルダがありません: {d}"}, ensure_ascii=False))
        return 1
    fn = {"tabs": cmd_tabs, "read": cmd_read, "append": cmd_append, "update": cmd_update,
          "update-status": cmd_update_status, "next-id": cmd_next_id, "expand-schedule": cmd_expand}[a.cmd]
    try:
        out = fn(d, a)
    except (LedgerError, json.JSONDecodeError, KeyError) as e:
        print(json.dumps({"error": str(e)}, ensure_ascii=False))
        return 1
    code = 0
    if isinstance(out, tuple):
        out, code = out
    print(json.dumps(out, ensure_ascii=False, indent=1))
    return code


if __name__ == "__main__":
    sys.exit(main())
