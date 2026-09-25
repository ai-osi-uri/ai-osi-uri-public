#!/usr/bin/env python3
"""OSI Finance 台帳テンプレートを data-layout.yaml から生成・照合する

使い方:
    python3 build_templates.py                      # assets/templates/ に4本を作り直す
    python3 build_templates.py --check              # assets/templates/ の4本を yaml と照合（差分があれば exit 1）
    python3 build_templates.py --ledger-set <DIR>   # <DIR> に空の台帳一式（請求管理台帳.xlsx など）を置く。既存は上書きしない
    python3 build_templates.py --check --ledgers <DIR> [--allow-data]
                                                    # 実台帳（名前はテンプレートでなく台帳名）を照合する

なぜこれがあるか
----------------
テンプレートを手で直していたら、仕様（data-layout.yaml）と列・見出し行・タブがずれたまま配られ、
別環境で手順どおりに動かなかった（契約マスタに列が足りない／廃止列が残る／タブが無い／記入例の行が
月初の請求処理に拾われる）。テンプレートは yaml から機械的に作り、照合も機械的に行う。

テンプレートの約束
------------------
* データ行は空。記入例は「00_使い方_◯◯」タブの説明にだけ書く（台帳の行に置くと処理が拾う）。
* 発行者設定の値は空欄（既定値のある項目だけ既定値）。setup が settings から埋める（fill_issuer_settings.py）。
* 勘定科目マスタだけは標準科目を入れておく（無いと BS が成立しない。data-layout.yaml の note）。
* optional: true のタブ（消化記録）は作らない。使う組織が足す。
"""
import argparse
import shutil
import sys
from pathlib import Path

try:
    import openpyxl
    import yaml
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
except ImportError as e:
    sys.exit(f"openpyxl と PyYAML が必要です: pip3 install openpyxl pyyaml ({e})")

HERE = Path(__file__).resolve().parent
PLUGIN = HERE.parent.parent
SCHEMA = PLUGIN / "assets" / "schema" / "data-layout.yaml"
TEMPLATES = PLUGIN / "assets" / "templates"

ORDER = ["共通マスタ", "請求管理台帳", "支払管理台帳", "仕訳台帳"]


def template_name(ledger):
    return f"{ledger}_テンプレート.xlsx"


def usage_tab(ledger):
    # タブ名は台帳をまたいで一意でなければならない（sheets_* はタブ名だけで引く）
    return f"00_使い方_{ledger}"


# 勘定科目マスタの標準科目（data-layout.yaml 勘定科目マスタ note の「標準に必ず含めること」＋スキルが使う科目）
# (科目コード, 科目名, 区分, 税区分既定, MF科目名)
STD_ACCOUNTS = [
    ("1000", "現金", "資産", "対象外", "現金"),
    ("1010", "普通預金", "資産", "対象外", "普通預金"),
    ("1020", "売掛金", "資産", "対象外", "売掛金"),
    ("1030", "仮払消費税", "資産", "対象外", "仮払消費税"),
    ("1040", "前払費用", "資産", "対象外", "前払費用"),
    ("1050", "前払金", "資産", "対象外", "前払金"),
    ("1060", "短期貸付金", "資産", "対象外", "短期貸付金"),
    ("1500", "差入保証金", "資産", "対象外", "差入保証金"),
    ("1510", "敷金", "資産", "対象外", "敷金"),
    ("1520", "長期前払費用", "資産", "対象外", "長期前払費用"),
    ("2010", "未払金", "負債", "対象外", "未払金"),
    ("2020", "預り金", "負債", "対象外", "預り金"),
    ("2030", "仮受消費税", "負債", "対象外", "仮受消費税"),
    ("2040", "買掛金", "負債", "対象外", "買掛金"),
    ("2050", "前受金", "負債", "対象外", "前受金"),
    ("2060", "未払費用", "負債", "対象外", "未払費用"),
    ("2070", "短期借入金", "負債", "対象外", "短期借入金"),
    ("2500", "長期借入金", "負債", "対象外", "長期借入金"),
    ("3010", "資本金", "純資産", "対象外", "資本金"),
    ("3020", "繰越利益剰余金", "純資産", "対象外", "繰越利益剰余金"),
    ("4010", "売上高", "収益", "課税売上10%", "売上高"),
    ("4020", "雑収入", "収益", "対象外", "雑収入"),
    ("5010", "業務委託費", "費用", "課税仕入10%", "業務委託料"),
    ("5020", "外注費", "費用", "課税仕入10%", "外注費"),
    ("5030", "支払報酬料", "費用", "課税仕入10%", "支払報酬"),
    ("5040", "地代家賃", "費用", "課税仕入10%", "地代家賃"),
    ("5050", "通信費", "費用", "課税仕入10%", "通信費"),
    ("5060", "旅費交通費", "費用", "課税仕入10%", "旅費交通費"),
    ("5070", "交際費", "費用", "課税仕入10%", "接待交際費"),
    ("5080", "会議費", "費用", "課税仕入10%", "会議費"),
    ("5090", "消耗品費", "費用", "課税仕入10%", "備品・消耗品費"),
    ("5100", "広告宣伝費", "費用", "課税仕入10%", "広告宣伝費"),
    ("5110", "支払手数料", "費用", "課税仕入10%", "支払手数料"),
    ("5120", "新聞図書費", "費用", "課税仕入10%", "新聞図書費"),
    ("5130", "保険料", "費用", "非課税仕入", "保険料"),
    ("5140", "租税公課", "費用", "対象外", "租税公課"),
    ("5150", "減価償却費", "費用", "対象外", "減価償却費"),
    ("5160", "支払利息", "費用", "非課税仕入", "支払利息"),
]

# 使い方タブの本文。固有名詞は書かない（例は サンプル株式会社 / 山田太郎 / example.com）
COMMON_HEAD = [
    "・置き場所：経理フォルダ（osi-profile.md の paths.finance ＝ 拡張の設定「台帳フォルダ」xlsx_ledger_dir）。"
    "共通マスタ・請求管理台帳・支払管理台帳・仕訳台帳・コンソールを開く.html を同じフォルダに置く。",
    "・データ行は空です。記入例は下にだけ書いてあります（台帳の行に例を置くと、月初の請求処理などが本物として拾います）。",
    "・列はヘッダー名で読みます。列の追加・並べ替えはかまいませんが、列名は変えないでください。",
    "・分からない値は空欄にして、理由を備考に書きます（「要確認」などの文字で埋めない）。",
    "・仕様の正本：プラグインの assets/schema/data-layout.yaml。テンプレートは assets/scripts/build_templates.py で生成。",
]
USAGE = {
    "共通マスタ": [
        "共通マスタ テンプレート（取引先マスタ）  ※OSI Finance",
        "・ファイル名は『共通マスタ.xlsx』にして置きます。",
        *COMMON_HEAD,
        "・取引先マスタ＝会社そのもの（1社1行）。受注も発注も同じ1行にぶら下げます。正式名称・住所・請求書の宛先はここだけが持ちます。",
        "・取引先ID は P-001 形式の文字列（裸の数字にしない）。契約取込（osi-finance-contract-intake）が採番して足します。",
        "",
        "記入例（参考。台帳には書かない）",
        "  取引先ID=P-001 / 正式名称=サンプル株式会社 / 略称=サンプル / 区分=顧客 / 郵便番号=100-0001 / 住所=東京都千代田区… /"
        " 請求先To=山田太郎 / Toアドレス=yamada@example.com",
    ],
    "請求管理台帳": [
        "請求管理台帳 テンプレート（AR・もらう側）  ※OSI Finance",
        "・ファイル名は『請求管理台帳.xlsx』（または settings の LEDGER_BILLING_FILE）にして置きます。",
        *COMMON_HEAD,
        "・発行者設定：値は osi-finance-setup が osi-finance-settings.md から埋めます（assets/scripts/fill_issuer_settings.py）。"
        "空欄のままでは請求書を発行しません（osi-finance-invoice が止めて知らせます）。settings を変えたら埋め直します。",
        "・契約マスタ：1契約1行。契約ID=AR-C-001（受注）／AP-C-001（発注）。会社の名前・住所・宛先は持たず、"
        "取引先ID で共通マスタの取引先マスタを引きます。状態=締結済 の行だけが請求予定に展開されます。",
        "・雛形マスタ：契約書の雛形（TPL-001…）。osi-finance-contract-draft が登録します。",
        "・月次請求スケジュール：1行=1請求。osi-finance-contract-intake が契約から展開し、"
        "請求書番号（INV-YYYY-MM-連番3桁、YYYY-MM=請求月＝請求日の月）は osi-finance-invoice が発行するときに振ります。"
        "「今月の請求書」は請求月が今月の未請求行です。発行元＝MF／freee／ローカル。"
        "下書き済＝請求書PDFを作った・まだ送っていない（メール下書きは任意）。",
        "・拡張（.mcpb）が無い環境では assets/scripts/ledger_io.py で読み書きします（タブ名＋列名で指定）。",
        "・前受・消化型（回数券など）を使う組織だけ「消化記録」タブを足します（settings §4-4）。",
        "",
        "記入例（参考。台帳には書かない）",
        "  契約マスタ：契約ID=AR-C-001 / 取引先=サンプル / 契約種別=業務委託(準委任) / 金額区分=月額固定 / 月額(税込)=220000 /"
        " 税率=10% / 契約開始=2026-04-01 / 契約終了=2027-03-31 / 取引先ID=P-001 / 方向=受注 / 状態=締結済 / 出所=手動",
        "  月次請求スケジュール：契約ID=AR-C-001 / 対象月=2026-04 / 件名=業務委託（2026年4月分） / 請求額(税込)=220000 /"
        " 請求日=2026-05-01 / 支払期限=2026-05-31 / 請求ステータス=未請求（この行は5月に「今月の請求書」で拾われ INV-2026-05-… になる）",
    ],
    "支払管理台帳": [
        "支払管理台帳 テンプレート（AP・払う側）  ※OSI Finance",
        "・ファイル名は『支払管理台帳.xlsx』（または settings の LEDGER_PAYMENT_FILE）にして置きます。",
        *COMMON_HEAD,
        "・見出し行の位置は仕様どおりです：支払先マスタ・経費一覧・源泉徴収管理は3行目、月次支払管理は4行目、月次支払スケジュールは1行目。",
        "・支払先マスタ：振込先口座・源泉の要否・支払サイクル。会社の名前は 取引先ID で共通マスタの取引先マスタを引きます。"
        "既定の支払方法で 振込／カード を振り分けます（カード払いは支払予定に起票しない＝二重払い防止）。",
        "・受領請求書を月次支払管理に起票するときは 受領請求書ファイル（01.受領請求書/YYYY-MM/…pdf）を必ず書きます。",
        "・月次支払管理＝支払の実績。月次支払スケジュール＝契約から展開した支払の予定。経費一覧＝明細の蓄積（ここから仕訳を起こさない）。",
        "・源泉徴収管理：源泉対象の支払の税額と納付状況。",
        "",
        "記入例（参考。台帳には書かない）",
        "  支払先マスタ：支払先ID=V-001 / 支払カテゴリ=外注費 / 取引先ID=P-001 /"
        " 既定の支払方法=振込 / 源泉徴収対象=なし / ステータス=支払対象",
    ],
    "仕訳台帳": [
        "仕訳台帳 テンプレート（内部仕訳帳）  ※OSI Finance",
        "・ファイル名は『仕訳台帳.xlsx』（または settings の LEDGER_JOURNAL_FILE）にして置きます。",
        *COMMON_HEAD,
        "・仕訳帳：osi-finance-journal が起こします。記帳の前後で assets/scripts/verify_ledger.py を走らせ、ERROR があれば記帳しません。",
        "・勘定科目マスタ：標準科目を入れてあります（これは記入例ではなく使う値）。MF科目名／freee科目名 を自社の会計SaaSに合わせます。",
        "・月次サマリ：会計SaaS の損益の写し。ACCOUNTING_SYNC=none の組織は空のままでよい。",
        "・会計SaaS を使わない組織は、会計年度の開始日に 出所=期首残高 の仕訳を入れます（無いと BS が成立しません）。",
    ],
}

HEADER_FONT = Font(bold=True)
HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")
TEXT_COL_HINTS = ("ID", "番号", "対象月", "発生月", "発生日", "郵便番号", "コード", "旧ID", "envelope_id")


def load_schema():
    return yaml.safe_load(SCHEMA.read_text(encoding="utf-8"))


def tabs_of(schema, ledger, include_optional=False):
    for tab, spec in (schema["ledgers"][ledger]["tabs"] or {}).items():
        if spec.get("optional") and not include_optional:
            continue
        yield tab, spec


def col_names(spec):
    return [c["name"] for c in spec.get("columns", [])]


def build_book(schema, ledger):
    wb = openpyxl.Workbook()
    ws0 = wb.active
    ws0.title = usage_tab(ledger)
    for i, line in enumerate(USAGE[ledger], start=1):
        ws0.cell(row=i, column=1, value=line or None)
    ws0["A1"].font = Font(bold=True, size=12)
    ws0.column_dimensions["A"].width = 140
    for row in ws0.iter_rows():
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")

    for tab, spec in tabs_of(schema, ledger):
        ws = wb.create_sheet(tab)
        hr = int(spec.get("header_row", 1))
        cols = col_names(spec)
        if hr == 3:
            # 仕様：1行目が空、2行目が凡例（header.js は非空セルが最も多い行を見出しとみなす）
            ws.cell(row=2, column=1, value=f"凡例：{tab}。見出しは3行目。列の意味は「{usage_tab(ledger)}」タブを参照")
        for j, name in enumerate(cols, start=1):
            c = ws.cell(row=hr, column=j, value=name)
            c.font = HEADER_FONT
            c.fill = HEADER_FILL
            letter = get_column_letter(j)
            ws.column_dimensions[letter].width = max(10, min(40, len(name) * 2 + 4))
            if any(h in name for h in TEXT_COL_HINTS):
                ws.column_dimensions[letter].number_format = "@"  # 001 や 2026-04 を Excel に数値・日付化させない
            vals = next((c2.get("values") for c2 in spec["columns"] if c2["name"] == name), None)
            if vals:
                dv = DataValidation(type="list", formula1='"' + ",".join(vals) + '"', allow_blank=True,
                                    errorStyle="warning", showErrorMessage=True,
                                    errorTitle="仕様外の値", error="data-layout.yaml の values 以外の値です")
                dv.add(f"{letter}{hr + 1}:{letter}2000")
                ws.add_data_validation(dv)
        ws.freeze_panes = f"A{hr + 1}"  # 文字列で渡す（cell() だと空セルが生まれ、データ行があるように見える）

        if tab == "発行者設定":
            for i, it in enumerate(spec.get("items", []), start=hr + 1):
                ws.cell(row=i, column=1, value=it["name"])
                ws.cell(row=i, column=2, value=it.get("default"))
                note = f"settings の {it['settings']} から setup が記入"
                if it.get("note"):
                    note = f"{it['note']}。{note}"
                ws.cell(row=i, column=3, value=note)
            ws.column_dimensions["B"].width = 44
            ws.column_dimensions["C"].width = 60
        if tab == "勘定科目マスタ":
            idx = {n: k for k, n in enumerate(cols)}
            for i, acc in enumerate(STD_ACCOUNTS, start=hr + 1):
                code, name, kubun, tax, mf = acc
                for key, v in (("科目コード", code), ("科目名", name), ("区分", kubun),
                               ("税区分既定", tax), ("MF科目名", mf)):
                    ws.cell(row=i, column=idx[key] + 1, value=v)
    return wb


def write_templates(schema, outdir):
    outdir.mkdir(parents=True, exist_ok=True)
    for ledger in ORDER:
        p = outdir / template_name(ledger)
        build_book(schema, ledger).save(p)
        print(f"作成: {p}")


def ledger_set(schema, outdir, names=None):
    outdir.mkdir(parents=True, exist_ok=True)
    names = names or {}
    for ledger in ORDER:
        src = TEMPLATES / template_name(ledger)
        dst = outdir / f"{names.get(ledger, ledger)}.xlsx"
        if dst.exists():
            print(f"既存のため上書きしない: {dst}")
            continue
        shutil.copyfile(src, dst)
        print(f"配置: {dst}")


# ---------------------------------------------------------------- 照合

def row_values(ws, r):
    vals = [c.value for c in ws[r]] if r <= ws.max_row else []
    vals = [str(v).strip() if v is not None else "" for v in vals]
    while vals and vals[-1] == "":
        vals.pop()
    return vals


def check(schema, d, template_mode=True, allow_data=False):
    errs, warns = [], []
    seen_tabs = {}
    for ledger in ORDER:
        p = d / (template_name(ledger) if template_mode else f"{ledger}.xlsx")
        if not p.exists():
            errs.append(f"{p.name}: ファイルが無い")
            continue
        wb = openpyxl.load_workbook(p)
        expected = {t for t, _ in tabs_of(schema, ledger)}
        optional = {t for t, s in tabs_of(schema, ledger, True) if s.get("optional")}
        for name in wb.sheetnames:
            seen_tabs.setdefault(name, []).append(p.name)
            if name not in expected | optional and not name.startswith("00_使い方"):
                warns.append(f"{p.name}: 仕様に無いタブ「{name}」")
        for tab, spec in tabs_of(schema, ledger, True):
            if tab not in wb.sheetnames:
                if not spec.get("optional"):
                    errs.append(f"{p.name}: タブ「{tab}」が無い")
                continue
            ws = wb[tab]
            hr = int(spec.get("header_row", 1))
            got = row_values(ws, hr)
            want = col_names(spec)
            if got != want:
                missing = [c for c in want if c not in got]
                extra = [c for c in got if c and c not in want]
                errs.append(f"{p.name}/{tab}: 見出し({hr}行目)が仕様と不一致"
                            f"{' 不足=' + str(missing) if missing else ''}{' 余分=' + str(extra) if extra else ''}"
                            f"{' 並び違い' if not missing and not extra else ''}")
            for r in range(1, hr):
                if len([v for v in row_values(ws, r) if v]) >= len(want) and want:
                    errs.append(f"{p.name}/{tab}: {r}行目に見出しより列の多い行がある（見出し行の誤検出の恐れ）")
            if tab == "発行者設定":
                items = [it["name"] for it in spec.get("items", [])]
                keys = [str(ws.cell(row=r, column=1).value or "").strip() for r in range(hr + 1, ws.max_row + 1)]
                miss = [k for k in items if k not in keys]
                if miss:
                    (errs if template_mode else warns).append(f"{p.name}/発行者設定: 項目が無い {miss}")
                if template_mode:
                    for r in range(hr + 1, ws.max_row + 1):
                        k = str(ws.cell(row=r, column=1).value or "").strip()
                        v = ws.cell(row=r, column=2).value
                        it = next((x for x in spec.get("items", []) if x["name"] == k), None)
                        if it and (v or None) != it.get("default"):
                            errs.append(f"{p.name}/発行者設定: 「{k}」の値がテンプレートの既定と違う（{v!r}）")
                continue
            if tab == "勘定科目マスタ" or allow_data:
                continue
            filled = [r for r in range(hr + 1, ws.max_row + 1) if any(v for v in row_values(ws, r))]
            if filled:
                errs.append(f"{p.name}/{tab}: データ行が空でない（{len(filled)}行。記入例が残っている）")
    for tab, files in seen_tabs.items():
        if len(files) > 1:
            errs.append(f"タブ名「{tab}」が複数のファイルにある: {files}（sheets_* はタブ名で引くため一意にする）")
    return errs, warns


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--check", action="store_true", help="作らずに照合だけする")
    ap.add_argument("--out", type=Path, default=TEMPLATES, help="テンプレートの出力先（既定: assets/templates）")
    ap.add_argument("--ledgers", type=Path, help="--check で実台帳（台帳名.xlsx）のフォルダを照合する")
    ap.add_argument("--allow-data", action="store_true", help="--check でデータ行があっても咎めない（実台帳用）")
    ap.add_argument("--ledger-set", type=Path, help="空の台帳一式をこのフォルダに置く（既存は上書きしない）")
    a = ap.parse_args()
    schema = load_schema()

    if a.ledger_set:
        ledger_set(schema, a.ledger_set)
        return 0
    if a.check:
        d = a.ledgers or a.out
        errs, warns = check(schema, d, template_mode=a.ledgers is None, allow_data=a.allow_data)
        for w in warns:
            print(f"! {w}")
        for e in errs:
            print(f"✗ {e}")
        n_tabs = sum(1 for lg in ORDER for _ in tabs_of(schema, lg))
        print(f"{'OK' if not errs else 'NG'}: {d} を data-layout.yaml（layout_version {schema.get('layout_version')}）と照合"
              f" — 台帳 {len(ORDER)} / タブ {n_tabs} / ERROR {len(errs)} / WARN {len(warns)}")
        return 1 if errs else 0
    write_templates(schema, a.out)
    return 0


if __name__ == "__main__":
    sys.exit(main())
