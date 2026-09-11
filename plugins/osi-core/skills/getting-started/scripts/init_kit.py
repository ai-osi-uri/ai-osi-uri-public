#!/usr/bin/env python3
"""init_kit.py — 会社プロファイル（osi-profile.md）とプロジェクトキットを、連結フォルダに作る。

  python3 init_kit.py --root <連結フォルダ> --answers answers.json [--dry]
  python3 init_kit.py --root <連結フォルダ> --status          # 何があって何が無いかだけ出す

answers.json（getting-started スキルが 5 問の回答から組む）:
  {
    "company": {"name": "サンプル株式会社", "name_display": "Sample Inc.", "representative": "代表取締役 山田 太郎",
                "address": "...", "domain": "example.com", "url": "https://example.com/", "contact_email": "info@example.com",
                "business": "...", "slug": "sample", "reverse_domain": "com.example"},
    "members": {"owner": "山田"},
    "layout": "new" | "existing",              # new: 規定の雛形フォルダを作る / existing: paths を既存に合わせる
    "paths": {"projects": "案件", ...},        # existing のときの上書き（root からの相対）
    "ledgers": {"sales": "営業管理/営業管理表.xlsx", ...},
    "connectors": {"deploy": false, "finance": false, ...},
    "sample_data": true                        # ダミー案件 1 件と台帳 1 行を入れる（初回の動作確認用）
  }

不変条件:
  - 既にあるファイルは上書きしない（osi-profile.md も台帳も）。無いものだけ作る
  - 機微値（登録番号・法人番号・口座）は answers に無ければ空欄で作る。チャットで聞かない
  - 出力は最後に JSON 1 行（created / skipped / profile）。人が読む説明はスキル側が付ける
"""
import argparse, json, os, re, sys, datetime

HERE = os.path.dirname(os.path.abspath(__file__))
TEMPLATE = os.path.join(HERE, "..", "assets", "osi-profile.example.md")

SALES_COLUMNS = ["案件ID", "取引先名", "URL", "事業区分", "紹介元", "獲得日", "ステータス", "確度", "営業担当",
                 "CAIO担当", "サポート", "アサイン状態", "次アクション", "次アクション期日", "最終接触日",
                 "想定額(税込)", "課金形態", "想定開始月", "想定期間(ヶ月)", "窓口氏名", "窓口メール",
                 "議事録", "提案書", "失注-保留日", "失注-保留理由", "再アプローチ月", "顧客課題メモ"]
STATUSES = ["リード", "商談中", "提案中", "金額提示済み", "受注", "契約済み", "デリバリー中", "保留", "失注", "終了"]
SUBFOLDERS = ["01_提案・見積", "02_契約", "03_制作・成果物", "04_請求", "05_受領資料"]

DEFAULT_PATHS = {"root": ".", "projects": "案件", "project_naming": "{案件ID3桁}.{企業名}",
                 "project_template": "案件/000.フォルダ構造テンプレート", "shared": "共通", "sales_admin": "営業管理",
                 "sales_materials": "営業資料", "finance": "経理", "contracts": "経理/00.契約書", "billing": "経理",
                 "expenses": "経理/03.経費管理", "pr_articles": "広報/案件記事", "vault": ""}
DEFAULT_LEDGERS = {"sales": "営業管理/営業管理表.xlsx", "sales_tab": "取引先管理", "sales_header_row": 5,
                   "sales_rules": "営業管理/営業管理表_運用ルール.md", "apps": "案件/_アプリ台帳.md"}
CONNECTOR_KEYS = ["deploy", "finance", "creative", "obsidian", "money_forward", "docusign", "plaud", "slack", "gmail", "google_calendar"]


def render_profile(ans):
    """雛形（assets/osi-profile.example.md）の `{{ }}` を回答で埋める。無いキーは空文字。"""
    tpl = open(TEMPLATE, encoding="utf-8").read()
    m = re.match(r"^---\n(.*?)\n---\n(.*)$", tpl, re.S)
    fm, body = m.group(1), m.group(2)
    company = ans.get("company", {}) or {}
    members = ans.get("members", {}) or {}
    paths = dict(DEFAULT_PATHS); paths.update(ans.get("paths", {}) or {})
    ledgers = dict(DEFAULT_LEDGERS); ledgers.update(ans.get("ledgers", {}) or {})
    conns = {k: bool((ans.get("connectors", {}) or {}).get(k, False)) for k in CONNECTOR_KEYS}
    vals = {"company": company, "members": members, "paths": paths, "ledgers": ledgers, "connectors": conns}
    out, section = [], None
    for line in fm.split("\n"):
        mm = re.match(r"^(\w+):\s*$", line)
        if mm:
            section = mm.group(1); out.append(line); continue
        km = re.match(r'^(  )(\w+):\s*(.*)$', line)
        if km and section in vals:
            key, rest = km.group(2), km.group(3)
            comment = ""
            cm = re.search(r"\s+#\s.*$", rest)
            if cm: comment = cm.group(0)
            v = vals[section].get(key)
            if section == "connectors":
                out.append(f"  {key}: {'true' if v else 'false'}{comment}"); continue
            if key == "team":
                team = members.get("team") or []
                out.append("  team: " + json.dumps(team, ensure_ascii=False) + comment); continue
            if key == "project_subfolders":
                out.append("  project_subfolders: " + json.dumps(SUBFOLDERS, ensure_ascii=False)); continue
            if isinstance(v, bool): sv = "true" if v else "false"
            elif isinstance(v, (int, float)): sv = str(v)
            else: sv = json.dumps("" if v is None else str(v), ensure_ascii=False)
            out.append(f"  {key}: {sv}{comment}"); continue
        if line.startswith("# このファイルを") or line.startswith("# 実値版はコミット"):
            continue
        out.append(line)
    fm2 = "\n".join(out).replace("# osi-profile — 会社プロファイル（雛形）", f"# osi-profile — {company.get('name_display') or company.get('name') or '会社'} の会社プロファイル（実値版・Git には入れない）")
    note = f"\n# {company.get('name_display') or company.get('name') or '会社'} のプロファイル\n\n- 生成: getting-started（{datetime.date.today().isoformat()}）。値の修正はこのファイルを直接編集する。\n- 経理の詳細設定は `{paths['finance']}/osi-finance-settings.md`（osi-finance-setup が作る）。\n"
    return "---\n" + fm2 + "\n---\n" + note, paths, ledgers


def make_ledger_xlsx(path, owner, sample):
    try:
        import openpyxl
    except ImportError:
        return "openpyxl が無いので台帳 xlsx は作れない（pip install openpyxl のあと再実行）"
    wb = openpyxl.Workbook(); ws = wb.active; ws.title = "取引先管理"
    ws.cell(1, 1, "営業管理表（案件マスタ）"); ws.cell(2, 1, "ヘッダは 5 行目。1 案件 1 行。行は動かさない。段階は「ステータス」列。")
    for c, h in enumerate(SALES_COLUMNS, 1): ws.cell(5, c, h)
    if sample:
        row = {"案件ID": "122", "取引先名": "サンプル株式会社", "URL": "https://example.com/", "事業区分": "AI伴走支援",
               "紹介元": "紹介者名", "獲得日": datetime.date.today().isoformat(), "ステータス": "商談中", "確度": "B",
               "営業担当": owner, "次アクション": "初回提案を作る", "窓口氏名": "山田 太郎", "窓口メール": "yamada@example.com",
               "顧客課題メモ": "問い合わせ対応が属人化している（ダミー）"}
        for c, h in enumerate(SALES_COLUMNS, 1): ws.cell(6, c, row.get(h, ""))
    ms = wb.create_sheet("マスタ"); ms.cell(1, 1, "ステータス")
    for i, s in enumerate(STATUSES, 2): ms.cell(i, 1, s)
    ms.cell(1, 2, "確度"); ms.cell(2, 2, "A"); ms.cell(3, 2, "B"); ms.cell(4, 2, "C")
    wb.save(path); return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--root", required=True); ap.add_argument("--answers"); ap.add_argument("--dry", action="store_true")
    ap.add_argument("--status", action="store_true")
    a = ap.parse_args()
    root = os.path.abspath(os.path.expanduser(a.root))
    profile_path = os.path.join(root, "osi-profile.md")
    if a.status:
        st = {"root": root, "root_exists": os.path.isdir(root), "profile": os.path.exists(profile_path),
              "shared_profile": os.path.exists(os.path.join(root, "_shared", "osi-profile.md")),
              "top_level": sorted(os.listdir(root))[:40] if os.path.isdir(root) else []}
        print(json.dumps(st, ensure_ascii=False)); return 0
    if not a.answers:
        print("answers.json が要る（--answers）", file=sys.stderr); return 2
    ans = json.load(open(a.answers, encoding="utf-8"))
    created, skipped, notes = [], [], []

    def ensure_dir(rel):
        p = os.path.join(root, rel)
        if os.path.isdir(p): skipped.append(rel + "/"); return p
        if not a.dry: os.makedirs(p, exist_ok=True)
        created.append(rel + "/"); return p

    def write_if_absent(rel, content):
        p = os.path.join(root, rel)
        if os.path.exists(p): skipped.append(rel); return
        if not a.dry:
            os.makedirs(os.path.dirname(p), exist_ok=True)
            open(p, "w", encoding="utf-8").write(content)
        created.append(rel)

    if not os.path.isdir(root):
        print(json.dumps({"error": f"連結フォルダが無い: {root}"}, ensure_ascii=False)); return 2
    text, paths, ledgers = render_profile(ans)
    if os.path.exists(profile_path):
        skipped.append("osi-profile.md"); notes.append("osi-profile.md は既にあるので触っていない。値を変えたいときは直接編集する")
    else:
        if not a.dry: open(profile_path, "w", encoding="utf-8").write(text)
        created.append("osi-profile.md")

    owner = (ans.get("members", {}) or {}).get("owner", "")
    layout = ans.get("layout", "new")
    if layout == "new" or ans.get("make_folders", True):
        for key in ("projects", "project_template", "shared", "sales_admin", "sales_materials", "finance", "pr_articles"):
            rel = paths.get(key)
            if rel: ensure_dir(rel)
        for sf in SUBFOLDERS: ensure_dir(os.path.join(paths["project_template"], sf))
        # 台帳
        lp = os.path.join(root, ledgers["sales"])
        if os.path.exists(lp): skipped.append(ledgers["sales"])
        else:
            if not a.dry:
                os.makedirs(os.path.dirname(lp), exist_ok=True)
                err = make_ledger_xlsx(lp, owner, ans.get("sample_data", True))
                if err: notes.append(err)
                else: created.append(ledgers["sales"])
            else: created.append(ledgers["sales"])
        write_if_absent(ledgers["sales_rules"], "# 営業管理表 運用ルール\n\n- 1 案件 1 行。行は動かさず、段階は「ステータス」列で表す（値は「マスタ」タブが正）。\n- 案件ID は 3 桁。採番の正本は案件フォルダ名（`{案件ID3桁}.{企業名}`）。台帳の案件ID 列に同じ値を入れる。\n- 書き込みは `osi-sales` の `ledger.py` から行う（バックアップを自動で取る）。\n")
        write_if_absent(ledgers["apps"], "# アプリ台帳\n\n| # | 状態 | アプリ名 | 案件 | 公開URL | リポ | 作成日 | 備考 |\n|---|---|---|---|---|---|---|---|\n")
        write_if_absent("CLAUDE.md", f"""# このフォルダの規約（getting-started が生成）

- 会社固有の値は `osi-profile.md`（このフォルダ直下）が正本。スキルはここから読む。
- 案件フォルダは `{paths['projects']}/{paths['project_naming']}/`。新規案件は `{paths['project_template']}/` と同じサブフォルダ構成で作る。
- 案件情報の正本は営業管理表 `{ledgers['sales']}`（タブ `{ledgers['sales_tab']}`・ヘッダ {ledgers['sales_header_row']} 行目）。
- 経理は `{paths['finance']}/`（詳細は osi-finance-setup が作る `osi-finance-settings.md`）。
- 成果物は案件フォルダの `03_制作・成果物/` に置く。このフォルダ直下に散らさない。
""")
        if ans.get("sample_data", True):
            sp = os.path.join(paths["projects"], "122.サンプル株式会社")
            for sf in SUBFOLDERS: ensure_dir(os.path.join(sp, sf))
            write_if_absent(os.path.join(sp, "05_受領資料", "初回ヒアリングメモ.md"),
                "# サンプル株式会社 初回ヒアリングメモ（ダミー）\n\n- 業種: 地域密着の小売（店舗 3・EC あり）。従業員 25 名\n- 困りごと: 問い合わせ対応が担当者 1 人に集中。返信が翌日以降になる\n- やりたいこと: よくある質問を先に返す仕組み。将来は在庫連動\n- 決裁: 社長。予算は月 10 万円程度から\n- 次回: 初回提案（試作案 3 つ）を持っていく\n")
    print(json.dumps({"profile": profile_path, "created": created, "skipped": skipped, "notes": notes, "dry": a.dry}, ensure_ascii=False))
    return 0


if __name__ == "__main__":
    sys.exit(main())
